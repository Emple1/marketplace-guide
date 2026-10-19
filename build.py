"""Сборка инструкции по модулю маркетплейсов.

Источник правды — общий Google Doc. Скрипт скачивает его в формате Word (.docx),
превращает в оформленный HTML и встраивает туда же .docx для кнопки «Скачать Word».

Запуск:
    python build.py                    # скачать Google Doc по ссылке из config.json
    python build.py путь/к/файлу.docx  # собрать из вручную скачанного .docx
    python build.py --separate-docx    # не встраивать .docx в HTML, положить рядом
    python build.py --no-compress      # не сжимать скриншоты
"""
import argparse
import base64
import csv
import datetime
import hashlib
import html
import io
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import mammoth
import requests
from bs4 import BeautifulSoup, NavigableString
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent
DIST = ROOT / "dist"
EMBED_LIMIT_MB = 40  # больше — .docx кладём рядом, иначе HTML станет неподъёмным
MAX_SIDE = 2000      # максимальная сторона скриншота, px
MAX_SHOW_HEIGHT = 900  # выше этого картинка показывается уменьшенной (по клику — полный размер)
WEB_DOCX = "instrukciya.docx"  # имя .docx на сайте (латиницей — надёжнее в адресе)
MEDIA_ATTR = re.compile(r'data-media="([^"]+)"')

sys.stdout.reconfigure(encoding="utf-8")

# Стили Google Docs / Word → HTML
STYLE_MAP = """
p[style-name='Title'] => h1.doc-title:fresh
p[style-name='title'] => h1.doc-title:fresh
p[style-name='Subtitle'] => p.doc-subtitle:fresh
p[style-name='subtitle'] => p.doc-subtitle:fresh
"""

# Абзац, начинающийся с «Слово:», превращается в цветную выноску
CALLOUTS = {
    "важно": "warn", "внимание": "warn", "обратите внимание": "warn",
    "осторожно": "danger", "нельзя": "danger",
    "совет": "tip", "подсказка": "tip", "лайфхак": "tip",
    "примечание": "info", "заметка": "info", "информация": "info",
}
CALLOUT_RE = re.compile(r"^\s*(" + "|".join(sorted(CALLOUTS, key=len, reverse=True)) + r")\s*[:!]", re.I)
ICONS = {
    "warn": '<path d="M12 9v4m0 4h.01M10.3 3.9 1.8 18a2 2 0 0 0 1.7 3h17a2 2 0 0 0 1.7-3L13.7 3.9a2 2 0 0 0-3.4 0z"/>',
    "danger": '<circle cx="12" cy="12" r="9"/><path d="m15 9-6 6m0-6 6 6"/>',
    "tip": '<path d="M9 18h6m-5 3h4M12 3a6 6 0 0 0-3.5 10.9c.6.4 1 1.1 1 1.8V16h5v-.3c0-.7.4-1.4 1-1.8A6 6 0 0 0 12 3z"/>',
    "info": '<circle cx="12" cy="12" r="9"/><path d="M12 11v5m0-8h.01"/>',
}
LINK_ICON = ('<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" '
             'stroke-linecap="round" stroke-linejoin="round"><path d="M10 13a5 5 0 0 0 7.5.5l3-3a5 5 0 0 0-7-7l-1.7 1.7"/>'
             '<path d="M14 11a5 5 0 0 0-7.5-.5l-3 3a5 5 0 0 0 7 7l1.7-1.7"/></svg>')
CAPTION_RE = re.compile(r"^\s*(рис\.|рисунок|скриншот)\s*\d*", re.I)

# Маркетплейсы для фильтра на странице: раздел относится к маркетплейсу, если тот назван
# в заголовке раздела (или выше по вложенности). Код → способы написания в заголовках.
MARKETPLACES = {
    "wb": r"wildberries|\bрвб\b|\bwb\b|\bвб\b",
    "ozon": r"ozon|озон",
    "ym": r"яндекс\s*\.?\s*маркет|\bям\b",
    "mm": r"мегамаркет|megamarket",
    "ali": r"aliexpress|алиэкспресс",
    "mvideo": r"м\.\s*видео|мвидео|эльдорадо",
    "lamoda": r"lamoda|ламода",
    "magnit": r"магнит\s*маркет",
}


def marketplaces_in(text):
    t = text.lower()
    return [code for code, rx in MARKETPLACES.items() if re.search(rx, t)]


def load_config():
    return json.loads((ROOT / "config.json").read_text(encoding="utf-8"))


def download_google_doc(doc):
    """Скачивает Google Doc как .docx. Нужен доступ «Все, у кого есть ссылка» (хотя бы чтение)."""
    m = re.search(r"/document/d/([\w-]+)", doc)
    doc_id = m.group(1) if m else doc.strip()
    url = f"https://docs.google.com/document/d/{doc_id}/export?format=docx"
    print(f"Скачиваю Google Doc {doc_id} …")
    r = requests.get(url, timeout=120)
    if r.status_code != 200 or "text/html" in r.headers.get("Content-Type", ""):
        sys.exit(
            "Не удалось скачать документ: у него нет доступа по ссылке.\n"
            "Либо включите в Google Doc «Настройки доступа → Все, у кого есть ссылка (читатель)»,\n"
            "либо скачайте вручную «Файл → Скачать → Microsoft Word (.docx)» и запустите:\n"
            "    python build.py путь/к/файлу.docx"
        )
    return r.content


def doc_id(doc):
    m = re.search(r"/document/d/([\w-]+)", doc)
    return m.group(1) if m else doc.strip()


def slugify(text, used):
    s = re.sub(r"[^\w\s-]", "", text.lower(), flags=re.U)
    s = re.sub(r"[\s_-]+", "-", s).strip("-")[:60].strip("-") or "razdel"
    base, n = s, 2
    while s in used:
        s, n = f"{base}-{n}", n + 1
    used.add(s)
    return s


def norm_space(t):
    return re.sub(r"\s+", " ", t).strip()


def text_of(tag):
    return re.sub(r"\s+", " ", tag.get_text(" ")).strip()


def postprocess(fragment, own_doc_id=None, support_url=None, collapsed=(), matrix_sheet=None):
    soup = BeautifulSoup(fragment, "html.parser")
    matrix = read_feature_matrix(soup, matrix_sheet)

    # Фоновые подложки страниц из документа (жёлтые ромбы на весь лист) — не скриншоты
    bg = soup.select("img[data-bg]")
    for img in bg:
        img.decompose()
    if bg:
        print(f"  · убрано фоновых подложек страниц: {len(bg)}")

    # Титул и подзаголовок документа показываются в шапке страницы, из текста убираем
    for t in soup.select("h1.doc-title, p.doc-subtitle"):
        t.decompose()
    # Титульный лист (всё до первого заголовка) — тоже, название уже в шапке
    first = soup.find(["h1", "h2", "h3"])
    if first is not None:
        for el in list(first.find_previous_siblings()):
            el.decompose()

    # Пустые абзацы, которые Google Docs оставляет между блоками
    for p in soup.find_all("p"):
        if not text_of(p) and not p.find("img"):
            p.decompose()

    # Ручное «Содержание» из документа не нужно — на странице есть боковое оглавление
    for h in soup.find_all("h1"):
        if text_of(h).lower() in ("содержание", "оглавление"):
            nxt = h.find_next_sibling()
            while nxt is not None and nxt.name != "h1":
                cur, nxt = nxt, nxt.find_next_sibling()
                cur.decompose()
            h.decompose()

    # «Заголовок», который целиком состоит из ссылок, — это пункты списка вопросов,
    # случайно оформленные стилем заголовка: делаем каждую ссылку обычным абзацем
    for h in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
        links = h.find_all("a", href=True)
        if links and norm_space(" ".join(text_of(a) for a in links)) == text_of(h):
            for a in links:
                p = soup.new_tag("p")
                p.append(a.extract())
                h.insert_before(p)
            for b in h.find_all("a", id=True):  # закладку заголовка переносим на первый абзац
                h.find_previous_sibling("p").insert(0, b.extract())
            h.decompose()

    # Заголовки: якоря для оглавления и ссылок «на раздел»
    used, aliases, orphans = set(), {}, []
    for h in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
        title = text_of(h)
        bookmarks = [a["id"] for a in h.find_all("a", href=False) if a.get("id")]  # закладки Google Docs
        if not title:
            orphans += bookmarks  # закладка пустого заголовка ведёт на следующий
            h.decompose()
            continue
        h["id"] = slugify(title, used)
        for b in bookmarks + orphans:
            aliases[b] = h["id"]
        orphans = []
        for a in h.find_all("a", href=False):
            a.unwrap()
        h["data-title"] = title
        anchor = soup.new_tag("a", attrs={"class": "anchor", "href": f"#{h['id']}",
                                          "title": "Скопировать ссылку на раздел"})
        anchor.append(BeautifulSoup(LINK_ICON, "html.parser"))
        h.append(anchor)

    # Внутренние ссылки документа («см. раздел …») ведут на новые якоря заголовков
    ids = {t["id"] for t in soup.find_all(id=True)}

    # Ссылки на этот же Google Doc («…/document/d/<id>/…#heading=h.xxx») уводили бы клиента
    # в исходник — превращаем их во внутренние ссылки на раздел инструкции
    if own_doc_id:
        for a in soup.find_all("a", href=lambda h: h and own_doc_id in h):
            m = re.search(r"#(?:heading=)?(h\.[\w-]+)", a["href"])
            if m:
                a["href"] = "#_heading=" + m.group(1)
                continue
            # Ссылка на документ целиком: убираем вместе с текстом-адресом
            parent = a.parent
            if text_of(a).startswith("http"):
                a.decompose()
            else:
                a.unwrap()
            if parent is not None and parent.name in ("p", "li") and len(text_of(parent)) < 80 and not parent.find("img"):
                parent.decompose()

    def norm(t):
        return re.sub(r"[\W_]+", " ", re.sub(r"\s\d+$", "", t)).strip().lower()

    def find_by_text(a):
        """Закладка потерялась — ищем раздел с тем же названием (заголовок или жирный абзац)."""
        want = norm(text_of(a))
        if not want:
            return None
        levels = ["h1", "h2", "h3", "h4", "h5", "h6"]
        below = [h for h in a.find_all_next(levels, limit=30) if h.get("id")]
        # 1) ближайший заголовок ниже с точно таким названием, 2) ниже — начинающийся с него
        #    («Wildberries» → «Wildberries - Настройка профиля»), 3) такой же в любом месте
        for h in below:
            if norm(h.get("data-title", "")) == want:
                return h["id"]
        for h in below:
            if norm(h.get("data-title", "")).startswith(want + " "):
                return h["id"]
        for h in soup.find_all(levels):
            if h.get("id") and norm(h.get("data-title", "")) == want:
                return h["id"]
        for p in soup.find_all("p"):
            if p is a.parent or p.find("a", href=True):
                continue
            st = p.find("strong")
            if st is not None and norm(text_of(p)) == want == norm(text_of(st)):
                if not p.get("id"):
                    p["id"] = slugify(text_of(p), used)
                    ids.add(p["id"])
                return p["id"]
        return None

    for a in soup.find_all("a", href=True):
        if not a["href"].startswith("#"):
            continue
        target = aliases.get(a["href"][1:], a["href"][1:])
        if target not in ids:
            target = find_by_text(a) or target
        if target in ids:
            a["href"] = "#" + target
        else:
            print(f"  · битая ссылка в документе: «{text_of(a)[:70]}» — раздел, на который она ведёт, удалён")
            a.unwrap()

    # Абзацы только из картинок → <figure>, следующая строка «Рис. …» → подпись
    for p in soup.find_all("p"):
        imgs = p.find_all("img")
        if not imgs or text_of(p):
            continue
        fig = soup.new_tag("figure")
        for img in imgs:
            img["loading"] = "lazy"
            if not img.get("alt"):
                img["alt"] = ""
            fig.append(img.extract())
        nxt = p.find_next_sibling()
        if nxt is not None and nxt.name == "p" and not nxt.find("img"):
            cap = text_of(nxt)
            whole_italic = nxt.find("em") is not None and text_of(nxt.find("em")) == cap
            if CAPTION_RE.match(cap) or (whole_italic and len(cap) < 200):
                fc = soup.new_tag("figcaption")
                for child in list(nxt.contents):
                    fc.append(child.extract())
                fig.append(fc)
                nxt.decompose()
        p.replace_with(fig)

    # Выноски «Важно: …», «Совет: …»
    for p in soup.find_all("p"):
        if p.find_parent(["table", "figure"]):
            continue
        m = CALLOUT_RE.match(text_of(p))
        if not m:
            continue
        kind = CALLOUTS[m.group(1).lower()]
        if p.find("strong") is None or not text_of(p).startswith(text_of(p.find("strong"))):
            # Выделяем слово-метку жирным, если автор этого не сделал
            first = p.contents[0] if p.contents else None
            if isinstance(first, NavigableString):
                mm = CALLOUT_RE.match(str(first))
                if mm:
                    strong = soup.new_tag("strong")
                    strong.string = str(first)[: mm.end()]
                    rest = str(first)[mm.end():]
                    first.replace_with(strong)
                    strong.insert_after(NavigableString(rest))
        box = soup.new_tag("div", attrs={"class": f"callout {kind}"})
        box.append(BeautifulSoup(
            f'<svg class="ico" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" '
            f'stroke-linecap="round" stroke-linejoin="round">{ICONS[kind]}</svg>', "html.parser"))
        body = soup.new_tag("div")
        p.replace_with(box)
        body.append(p)
        box.append(body)

    # Таблицы — с прокруткой на узком экране
    for table in soup.find_all("table"):
        if table.find_parent("table"):
            continue
        wrap = soup.new_tag("div", attrs={"class": "table-wrap"})
        table.wrap(wrap)

    # Ссылки наружу — в новой вкладке
    for a in soup.find_all("a", href=True):
        if a["href"].startswith("http"):
            a["target"] = "_blank"
            a["rel"] = "noopener"

    remove_marketplace_lists(soup)
    # Разделы из config.json → "collapsed" показываются свёрнутыми (кнопка «Показать полностью»)
    fold = {norm_space(t).lower() for t in collapsed}
    for h in soup.find_all(["h1", "h2", "h3"]):
        if h.get("data-title", "").lower() in fold:
            h["data-fold"] = "1"
    if support_url:
        link_support(soup, support_url)
    tag_marketplaces(soup, matrix)
    return str(soup)


# «обратитесь в службу поддержки», «заявки на техподдержку» … — ссылкой делаем слова про поддержку
SUPPORT_RE = re.compile(r"(?:обрат\w*\s+в\s+(?:нашу\s+)?|заявк\w*\s+на\s+)"
                        r"(службу\s+поддержки|техническую\s+поддержку|техподдержку|поддержку)", re.I)


def link_support(soup, url):
    """Упоминания обращения в техподдержку → ссылка на личный кабинет техподдержки."""
    count = 0
    for node in list(soup.find_all(string=SUPPORT_RE)):
        if node.find_parent("a") is not None:
            continue
        text, pos, parts = str(node), 0, []
        for m in SUPPORT_RE.finditer(text):
            parts.append(NavigableString(text[pos:m.start(1)]))
            a = soup.new_tag("a", href=url, target="_blank", rel="noopener")
            a.string = m.group(1)
            parts.append(a)
            pos = m.end(1)
            count += 1
        parts.append(NavigableString(text[pos:]))
        for part in parts:
            node.insert_before(part)
        node.extract()
    print(f"  · ссылок на техподдержку: {count}")


def is_marketplace_line(el):
    """Абзац, в котором только название маркетплейса («Wildberries», «OZON FBS»)."""
    if el is None or el.name != "p" or el.find("img") or len(text_of(el)) >= 40:
        return False
    own = marketplaces_in(text_of(el))
    if len(own) != 1:
        return False
    rest = re.sub(MARKETPLACES[own[0]], "", text_of(el).lower())
    return not re.sub(r"[\W_]|fbs|fbo", "", rest)


def remove_marketplace_lists(soup):
    """Убирает списки-перечисления «Wildberries / OZON / ЯндексМаркет …» (3+ строки подряд):
    это навигация из Google Doc, на странице её заменяют оглавление и фильтр по маркетплейсу."""
    removed = 0
    for el in soup.find_all("p", recursive=False):
        if el.parent is None or not is_marketplace_line(el) or is_marketplace_line(el.find_previous_sibling()):
            continue
        run = [el]
        while is_marketplace_line(run[-1].find_next_sibling()):
            run.append(run[-1].find_next_sibling())
        if len(run) >= 3:
            for x in run:
                x.decompose()
            removed += 1
    if removed:
        print(f"  · убрано списков-перечислений маркетплейсов: {removed}")


MATRIX_TITLE = "функционал по маркетплейсам"
MATRIX_YES = {"+", "✓", "✔", "v", "x", "х", "да", "есть", "?"}  # «?» = не уверены → считаем, что есть


def norm_title(t):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", t.lower().replace("ё", "е"))).strip()


MATRIX_FILE = ROOT / "matrix.csv"


def parse_matrix(rows):
    """Строки таблицы (списки текстов ячеек) → {название раздела: [коды маркетплейсов]}.

    Первая строка — шапка: «Раздел; WB; Ozon; …» (столбцы распознаются по названию маркетплейса),
    дальше — разделы как заголовки в документе (можно начало названия; «— » перед подразделом не мешает),
    в клетках «+» — есть, пусто — нет, «?» — не уверены (считаем, что есть)."""
    if not rows:
        return {}
    cols = [marketplaces_in(c) for c in rows[0]]
    matrix = {}
    for cells in rows[1:]:
        name = norm_title(cells[0]) if cells else ""
        if name:
            matrix[name] = [cols[i][0] for i, c in enumerate(cells[1:], 1)
                            if i < len(cols) and len(cols[i]) == 1 and c.strip().lower() in MATRIX_YES]
    return matrix


def load_matrix_file():
    """Запасная матрица функционала из matrix.csv в папке проекта (если Google Таблица недоступна)."""
    if not MATRIX_FILE.exists():
        return {}
    with open(MATRIX_FILE, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f, delimiter=";") if any(c.strip() for c in r)]
    matrix = parse_matrix(rows)
    if matrix:
        print(f"  · матрица функционала (matrix.csv): {len(matrix)} строк")
    return matrix


def load_matrix_sheet(url):
    """Матрица функционала из Google Таблицы (её заполняют коллеги; нужен доступ «все, у кого есть ссылка»)."""
    m = re.search(r"/spreadsheets/d/([\w-]+)", url)
    if not m:
        print(f"  · ВНИМАНИЕ: не понял ссылку на таблицу матрицы: {url}")
        return {}
    gid = re.search(r"[#&?]gid=(\d+)", url)
    # без gid Google отдаёт первый лист (у загруженных .xlsx номер листа не 0)
    export = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv" + (f"&gid={gid.group(1)}" if gid else "")
    try:
        r = requests.get(export, timeout=60)
        ok = r.status_code == 200 and "text/html" not in r.headers.get("Content-Type", "")
    except requests.RequestException:
        ok = False
    if not ok:
        print("  · ВНИМАНИЕ: не удалось скачать Google Таблицу матрицы (закрыт доступ по ссылке?) — беру matrix.csv")
        return {}
    rows = [row for row in csv.reader(io.StringIO(r.content.decode("utf-8-sig"))) if any(c.strip() for c in row)]
    matrix = parse_matrix(rows)
    print(f"  · матрица функционала (Google Таблица): {len(matrix)} строк")
    return matrix


def read_feature_matrix(soup, sheet_url=None):
    """Матрица функционала: Google Таблица (если указана в config.json → matrix_sheet), иначе matrix.csv;
    плюс запасной вариант — таблица «Функционал по маркетплейсам» в самом документе.

    Таблица из документа нужна только фильтру: её раздел убирается со страницы. Строки из документа
    дополняют/переопределяют остальные."""
    matrix = (load_matrix_sheet(sheet_url) if sheet_url else {}) or load_matrix_file()
    levels = ["h1", "h2", "h3", "h4", "h5", "h6"]
    head = next((h for h in soup.find_all(levels) if MATRIX_TITLE in text_of(h).lower()), None)
    if head is None:
        return matrix
    level, section, el = int(head.name[1]), [head], head.find_next_sibling()
    while el is not None and not (el.name in levels and int(el.name[1]) <= level):
        section.append(el)
        el = el.find_next_sibling()
    table = next((t for x in section for t in ([x] if x.name == "table" else x.find_all("table"))), None)
    if table is not None:
        doc_matrix = parse_matrix([[text_of(c) for c in tr.find_all(["td", "th"])] for tr in table.find_all("tr")])
        matrix.update(doc_matrix)
        print(f"  · матрица функционала из документа: {len(doc_matrix)} строк")
    for x in section:
        x.decompose()
    return matrix


def tag_marketplaces(soup, matrix=None):
    """Помечает блоки атрибутом data-mp="wb ozon …" — по ним работает фильтр на странице.

    Маркетплейсы раздела берутся из матрицы функционала (если раздел в ней есть), иначе из его
    заголовка («Wildberries - …»). Подразделы наследуют разметку раздела выше."""
    matrix = matrix or {}
    used_rows = set()

    def from_matrix(title):
        t = norm_title(title)
        # точное совпадение или строка матрицы — начало названия раздела (целыми словами)
        key = t if t in matrix else next((k for k in sorted(matrix, key=len, reverse=True)
                                          if t.startswith(k + " ")), None)
        if key is None:
            return None
        used_rows.add(key)
        return matrix[key] or ["none"]  # ни одной галочки — раздел не показывается ни при одном фильтре

    stack = []  # (уровень заголовка, маркетплейсы)
    for el in soup.find_all(recursive=False):
        if not hasattr(el, "name") or el.name is None:
            continue
        if re.fullmatch(r"h[1-6]", el.name):
            level = int(el.name[1])
            while stack and stack[-1][0] >= level:
                stack.pop()
            title = el.get("data-title", "")
            own = marketplaces_in(title)
            if own:  # маркетплейс назван в заголовке — по таким разделам строятся кнопки фильтра
                el["data-mp-own"] = " ".join(own)
            stack.append((level, own or from_matrix(title) or []))
        mps = next((m for _, m in reversed(stack) if m), [])
        # Строка из одного названия маркетплейса («Wildberries», «OZON» в списках)
        if not mps and el.name == "p" and len(text_of(el)) < 40:
            own = marketplaces_in(text_of(el))
            rest = re.sub(MARKETPLACES[own[0]], "", text_of(el).lower()) if len(own) == 1 else "x"
            if not re.sub(r"[\W_]|fbs|fbo", "", rest):  # кроме названия ничего нет
                mps = own
        if mps:
            el["data-mp"] = " ".join(mps)

    # Заголовок раздела показываем, если виден хоть один его подраздел (иначе подраздел
    # «повиснет» без родителя): добавляем заголовку маркетплейсы подразделов
    heads = [h for h in soup.find_all(re.compile(r"^h[1-6]$"), recursive=False) if h.get("data-mp")]
    for i, h in enumerate(heads):
        level, codes = int(h.name[1]), set(h["data-mp"].split())
        for sub in heads[i + 1:]:
            if int(sub.name[1]) <= level:
                break
            codes |= set(sub["data-mp"].split())
        codes.discard("none") if len(codes) > 1 else None
        h["data-mp"] = " ".join(sorted(codes))

    for key in matrix:
        if key not in used_rows:
            print(f"  · строка матрицы не нашла раздел: «{key}»")


def footer_html(cfg):
    text = html.escape(cfg.get("footer", ""))
    url = cfg.get("support_url")
    if url and "техподдержк" in text:
        text = text.replace("техподдержку", f'<a href="{html.escape(url)}" target="_blank" rel="noopener">техподдержку</a>', 1)
    return text


def logo_html(cfg):
    path = cfg.get("logo")
    if path and (ROOT / path).exists():
        p = ROOT / path
        mime = {"svg": "image/svg+xml", "jpg": "image/jpeg", "jpeg": "image/jpeg"}.get(p.suffix[1:].lower(), "image/png")
        data = p.read_bytes() if mime == "image/svg+xml" else transparent_corners(p.read_bytes())
        if data[1:4] == b"PNG":
            mime = "image/png"
        return f'<img src="data:{mime};base64,{base64.b64encode(data).decode()}" alt="">'
    letter = html.escape((cfg.get("title") or "И")[0])
    return f'<span class="badge-logo">{letter}</span>'


def transparent_corners(data, tolerance=30):
    """Белый фон вокруг логотипа → прозрачный (заливка от углов), чтобы в тёмной теме не было белого квадрата.

    Заливается только фон, связанный с краями картинки: белые места внутри логотипа остаются."""
    try:
        im = Image.open(io.BytesIO(data)).convert("RGBA")
    except Exception:
        return data
    w, h = im.size
    for xy in ((0, 0), (w - 1, 0), (0, h - 1), (w - 1, h - 1)):
        r, g, b, a = im.getpixel(xy)
        if a and min(r, g, b) >= 255 - tolerance:
            ImageDraw.floodfill(im, xy, (255, 255, 255, 0), thresh=tolerance)
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def web_image(data):
    """Скриншот для сайта: PNG → WebP без потерь, если так меньше (обычно на ~20 %). → (байты, расширение)."""
    try:
        im = Image.open(io.BytesIO(data))
        if im.format != "PNG":
            return data, None
        has_alpha = im.mode in ("RGBA", "LA") or "transparency" in im.info
        buf = io.BytesIO()
        im.convert("RGBA" if has_alpha else "RGB").save(buf, format="WEBP", lossless=True, method=4)
    except Exception:
        return data, None
    return (buf.getvalue(), ".webp") if buf.tell() < len(data) else (data, None)


def optimize_docx(docx_bytes):
    """Сжимает скриншоты внутри .docx без видимой потери качества.

    PNG → палитра 256 цветов (для скриншотов 1С разницы не видно, вес в 3–4 раза меньше),
    слишком большие картинки уменьшаются до MAX_SIDE. Размер картинок на странице Word
    задаётся в разметке документа, поэтому вёрстка Word не меняется.
    """
    src = zipfile.ZipFile(io.BytesIO(docx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename.startswith("word/media/") and info.filename.lower().endswith(".png"):
                try:
                    im = Image.open(io.BytesIO(data))
                    im = im.convert("RGBA") if im.mode in ("P", "LA", "RGBA", "PA") else im.convert("RGB")
                    if max(im.size) > MAX_SIDE:
                        im.thumbnail((MAX_SIDE, MAX_SIDE))
                    buf = io.BytesIO()
                    im.quantize(256, method=Image.Quantize.FASTOCTREE).save(buf, format="PNG", optimize=True)
                    if buf.tell() < len(data):
                        data = buf.getvalue()
                except Exception as e:  # битая картинка — оставляем как есть
                    print(f"  · не удалось сжать {info.filename}: {e}")
            dst.writestr(info, data)
    return out.getvalue()


def strip_matrix_from_docx(docx_bytes):
    """Убирает раздел «Функционал по маркетплейсам» из .docx для клиентов (он только для фильтра)."""
    import docx
    d = docx.Document(io.BytesIO(docx_bytes))

    def heading_level(p):
        m = re.search(r"(?:heading|заголовок)\s*(\d)", (p.style.name or "").lower()) if p.style is not None else None
        return int(m.group(1)) if m else None

    body = d.element.body
    paras = {p._p: p for p in d.paragraphs}
    start = level = None
    for p in d.paragraphs:
        lvl = heading_level(p)
        if lvl and MATRIX_TITLE in p.text.lower():
            start, level = p._p, lvl
            break
    if start is None:
        return docx_bytes
    remove, el = [start], start.getnext()
    while el is not None and el.tag != docx.oxml.ns.qn("w:sectPr"):
        p = paras.get(el)
        lvl = heading_level(p) if p is not None else None
        if lvl and lvl <= level:
            break
        remove.append(el)
        el = el.getnext()
    for el in remove:
        body.remove(el)
    out = io.BytesIO()
    d.save(out)
    print(f"  · матрица функционала убрана из Word-версии ({len(remove)} блоков)")
    return out.getvalue()


def background_media(docx_bytes):
    """Картинки, вставленные в документ «за текстом» (фоновые подложки страниц), — на сайте не нужны."""
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    doc = z.read("word/document.xml").decode("utf-8")
    rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
    targets = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'Id="([^"]+)"', m.group(0)), re.search(r'Target="([^"]+)"', m.group(0))
        if rid and tgt:
            targets[rid.group(1)] = "word/" + tgt.group(1).lstrip("/").removeprefix("word/")
    found = set()
    for attrs, body in re.findall(r"<wp:anchor\b([^>]*)>(.*?)</wp:anchor>", doc, re.S):
        if re.search(r'behindDoc="(1|true)"', attrs):
            found |= {targets[e] for e in re.findall(r'r:embed="([^"]+)"', body) if e in targets}
    return found


def media_index(docx_bytes):
    """Хэш содержимого картинки → её путь внутри .docx (word/media/imageN.png)."""
    z = zipfile.ZipFile(io.BytesIO(docx_bytes))
    return {hashlib.sha1(z.read(n)).hexdigest(): n for n in z.namelist() if n.startswith("word/media/")}


def main():
    ap = argparse.ArgumentParser(description="Сборка HTML + Word из Google Doc")
    ap.add_argument("docx", nargs="?", help="путь к .docx (иначе скачать Google Doc из config.json)")
    ap.add_argument("--separate-docx", action="store_true", help="положить .docx рядом, а не внутрь HTML")
    ap.add_argument("--no-compress", action="store_true", help="не сжимать скриншоты")
    args = ap.parse_args()
    cfg = load_config()

    if args.docx:
        docx_bytes = Path(args.docx).read_bytes()
    elif cfg.get("google_doc"):
        docx_bytes = download_google_doc(cfg["google_doc"])
    else:
        sys.exit("Укажите ссылку на Google Doc в config.json (поле google_doc) или передайте путь к .docx")

    if not args.no_compress:
        before = len(docx_bytes)
        print("Сжимаю скриншоты …")
        docx_bytes = optimize_docx(docx_bytes)
        print(f"  {before / 2**20:.1f} МБ → {len(docx_bytes) / 2**20:.1f} МБ")

    # Страница строится из полного документа (из него же читается матрица функционала),
    # а клиентам в Word уходит документ без матрицы
    source_bytes = docx_bytes
    docx_bytes = strip_matrix_from_docx(docx_bytes)
    size_mb = len(docx_bytes) / 2**20
    embed = not args.separate_docx and size_mb <= EMBED_LIMIT_MB

    # Картинки помечаем путём внутри .docx: в однофайловой версии браузер достаёт их из встроенного
    # .docx, в версии для сайта они лежат отдельными файлами
    media = media_index(docx_bytes)
    backgrounds = background_media(source_bytes)

    def convert_image(image):
        with image.open() as f:
            data = f.read()
        attrs = {"alt": image.alt_text or ""}
        try:  # размеры заранее — чтобы страница не «прыгала», пока грузятся картинки
            w, h = Image.open(io.BytesIO(data)).size
            if h > MAX_SHOW_HEIGHT:  # высокие блок-схемы — не выше экрана; полный размер — по клику
                w, h = max(1, round(w * MAX_SHOW_HEIGHT / h)), MAX_SHOW_HEIGHT
            attrs["width"], attrs["height"] = str(w), str(h)
        except Exception:
            pass
        path = media.get(hashlib.sha1(data).hexdigest())
        if path in backgrounds:  # фоновая подложка страницы — уберём при постобработке
            return {"data-bg": "1"}
        if path:
            attrs["data-media"] = path
        else:
            attrs["src"] = f"data:{image.content_type};base64,{base64.b64encode(data).decode()}"
        return attrs

    print("Превращаю в HTML …")
    result = mammoth.convert_to_html(io.BytesIO(source_bytes), style_map=STYLE_MAP,
                                     convert_image=mammoth.images.img_element(convert_image))
    for msg in result.messages:
        if "Unrecognised paragraph style" not in msg.message:
            print("  ·", msg.message)
    content = postprocess(result.value, doc_id(cfg.get("google_doc", "")), cfg.get("support_url"),
                          cfg.get("collapsed", []), cfg.get("matrix_sheet"))

    DIST.mkdir(exist_ok=True)
    name = cfg["file_name"]
    docx_name = f"{name}.docx"
    (DIST / docx_name).write_bytes(docx_bytes)
    zf = zipfile.ZipFile(io.BytesIO(docx_bytes))
    today = datetime.date.today().strftime("%d.%m.%Y")
    stats = {"headings": content.count("<h1") + content.count("<h2") + content.count("<h3"),
             "images": content.count("<img"), "built": today}

    # 1) Один файл — для отправки вручную: .docx встроен, картинки браузер достаёт из него
    if embed:
        single = content
        docx_href = "#"
        docx_data = ('<script type="application/octet-stream" id="docx-data">'
                     + base64.b64encode(docx_bytes).decode() + "</script>")
    else:
        def inline(m):
            data = zf.read(m.group(1))
            mime = Image.MIME.get(Image.open(io.BytesIO(data)).format, "image/png")
            return f'src="data:{mime};base64,{base64.b64encode(data).decode()}"'
        single = MEDIA_ATTR.sub(inline, content)
        docx_href, docx_data = docx_name, ""
    out = DIST / f"{name}.html"
    out.write_text(render_page(cfg, single, today, docx_name, docx_href, docx_data), encoding="utf-8")
    (DIST / "stats.json").write_text(json.dumps(stats, ensure_ascii=False), encoding="utf-8")

    # 2) Версия для сайта: лёгкая страница, скриншоты отдельными файлами (грузятся по мере
    #    прокрутки), Word скачивается по кнопке отдельным файлом
    web = DIST / "web"
    shutil.rmtree(web, ignore_errors=True)
    (web / "img").mkdir(parents=True)
    used = set(MEDIA_ATTR.findall(content))
    web_names = {}
    for path in used:
        data, ext = web_image(zf.read(path))
        web_names[path] = Path(path).stem + ext if ext else Path(path).name
        (web / "img" / web_names[path]).write_bytes(data)
    # data-src: картинку подставляет скрипт страницы заранее, за пару экранов до того, как до неё долистают
    web_content = MEDIA_ATTR.sub(lambda m: f'data-src="img/{web_names[m.group(1)]}" decoding="async"', content)
    web_content = web_content.replace(' loading="lazy"', '')
    (web / WEB_DOCX).write_bytes(docx_bytes)
    (web / "index.html").write_text(render_page(cfg, web_content, today, docx_name, WEB_DOCX, ""),
                                    encoding="utf-8")
    (web / "stats.json").write_text(json.dumps(stats, ensure_ascii=False), encoding="utf-8")
    (web / "sections.json").write_text(json.dumps(section_index(content, zf), ensure_ascii=False),
                                       encoding="utf-8")

    print(f"Готово: {out}  ({out.stat().st_size / 2**20:.1f} МБ)")
    print(f"        {DIST / docx_name}  ({size_mb:.1f} МБ)"
          + ("" if embed else "  ← отправляйте клиенту вместе с HTML"))
    print(f"        {web}  (для сайта: страница {(web / 'index.html').stat().st_size / 2**20:.1f} МБ"
          f" + {len(used)} картинок)")


def section_index(content, zf):
    """Отпечатки разделов для блока «Что нового»: [{id, title, level, hash}] по каждому h1–h3.

    hash — от собственного текста раздела (до следующего заголовка) и содержимого его картинок:
    перенумерация картинок и отметки фильтра (data-mp) изменением не считаются."""
    heads = list(re.finditer(r'<h([1-3])\b[^>]*\bid="([^"]+)"[^>]*>', content))
    index = []
    for i, m in enumerate(heads):
        chunk = content[m.end():heads[i + 1].start() if i + 1 < len(heads) else len(content)]
        imgs = [hashlib.sha1(zf.read(path)).hexdigest() for path in MEDIA_ATTR.findall(chunk)]
        text = norm_space(html.unescape(re.sub(r"<[^>]+>", " ", chunk)))
        title = re.search(r'data-title="([^"]*)"', m.group(0))
        index.append({"id": html.unescape(m.group(2)), "level": int(m.group(1)),
                      "title": html.unescape(title.group(1)) if title else text[:80],
                      "hash": hashlib.sha1((text + "|" + ",".join(imgs)).encode()).hexdigest()[:16]})
    return index


def render_page(cfg, content, today, docx_name, docx_href, docx_data):
    page = (ROOT / "template.html").read_text(encoding="utf-8")
    for key, value in {
        "TITLE": html.escape(cfg["title"]),
        "KICKER": html.escape(cfg.get("kicker", "")),
        "VERSION": html.escape(str(cfg.get("version", "1.0"))),
        "BUILD_DATE": today,
        "ACCENT": cfg.get("accent", "#2563eb"),
        "FOOTER": footer_html(cfg),
        "HOME_URL": html.escape(cfg.get("home_url") or "#"),
        "LOGO": logo_html(cfg),
        "DOCX_NAME": html.escape(docx_name),
        "DOCX_HREF": html.escape(docx_href),
    }.items():
        page = page.replace("{{" + key + "}}", value)
    # Содержимое подставляем последним, чтобы случайные {{…}} в тексте не трогать
    return page.replace("{{DOCX_DATA}}", docx_data).replace("{{CONTENT}}", content)

if __name__ == "__main__":
    main()
